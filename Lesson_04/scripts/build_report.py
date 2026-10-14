from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Результаты тестирования Lite.docx"


def set_cell_fill(cell, color):
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        borders.append(element)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{key}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_dialogue(doc, speaker, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.left_indent = Inches(0.18)
    run = p.add_run(speaker + "  ")
    run.bold = True
    run.font.color.rgb = RGBColor(24, 92, 60)
    p.add_run(text)


def add_result(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(13)
    run = p.add_run("Результат  ")
    run.bold = True
    run.font.color.rgb = RGBColor(24, 92, 60)
    p.add_run(text)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.7)
styles["Normal"].font.color.rgb = RGBColor(23, 33, 27)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(10)
title_ppr = styles["Title"].element.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
for name, size in (("Heading 1", 18), ("Heading 2", 13)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(14)
    style.paragraph_format.space_after = Pt(7)

title = doc.add_paragraph(style="Title")
title.add_run("Результаты тестирования AI куратора")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("Домашнее задание Lite по теме памяти контекста и работы с данными")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(86, 98, 90)

meta = doc.add_table(rows=3, cols=2)
meta.alignment = WD_TABLE_ALIGNMENT.LEFT
meta.autofit = False
meta.columns[0].width = Inches(1.55)
meta.columns[1].width = Inches(5.05)
for row, values in zip(meta.rows, [
    ("Приложение", "AI-куратор онлайн-школы Практика"),
    ("Дата проверки", "14 октября 2026 года"),
    ("Режим", "Локальный демо-режим без зависимости от внешнего API"),
]):
    for cell, value in zip(row.cells, values):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        set_cell_margins(cell)
    row.cells[0].paragraphs[0].runs[0].bold = True
    set_cell_fill(row.cells[0], "E8F2EC")

doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Итог проверки. ").bold = True
intro.add_run(
    "Все четыре обязательных сценария выполнены. Автоматический прогон завершился со статусом OK: "
    "6 тестов из 6. Долгосрочные сведения хранятся в локальном JSON, а контекст текущего диалога — "
    "в серверной сессии. SQL не использовался, поскольку он относится к заданию Pro."
)

h = doc.add_heading("Сводка результатов", level=1)
keep_with_next(h)
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = [Inches(0.62), Inches(2.25), Inches(2.85), Inches(0.82)]
headers = ["№", "Сценарий", "Проверяемое поведение", "Статус"]
for index, (cell, text, width) in enumerate(zip(table.rows[0].cells, headers, widths)):
    cell.text = text
    cell.width = width
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_fill(cell, "185C3C")
    set_cell_border(cell)
    set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
rows = [
    ("1", "Память текущего диалога", "Вторая реплика использует сведения из первой", "Пройден"),
    ("2", "Память после нового входа", "Профиль загружается после перезапуска приложения", "Пройден"),
    ("3", "Два разных клиента", "Ответы основаны на собственном профиле каждого клиента", "Пройден"),
    ("4", "Запрет чужих данных", "Клиент не получает административный доступ", "Пройден"),
]
for row_index, values in enumerate(rows):
    cells = table.add_row().cells
    for col, (cell, value, width) in enumerate(zip(cells, values, widths)):
        cell.text = value
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        set_cell_margins(cell)
        if row_index % 2:
            set_cell_fill(cell, "F3F7F4")
        if col in (0, 3):
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if col == 3:
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(24, 92, 60)
            cell.paragraphs[0].runs[0].bold = True

scenario_1 = doc.add_heading("Сценарий 1 Память текущего диалога", level=1)
scenario_1.paragraph_format.page_break_before = True
doc.add_paragraph(
    "Проверка подтверждает, что куратор получает историю активной сессии и не обрабатывает каждую реплику изолированно."
)
add_dialogue(doc, "Клиент", "Хочу понять основы циклов")
add_dialogue(doc, "Клиент", "А можешь продолжить рекомендацию?")
add_dialogue(doc, "Куратор", "Продолжаю текущий диалог и помню, что ранее вы написали: «Хочу понять основы циклов». Чем уточнить рекомендацию?")
add_result(doc, "Пройден. Вторая реплика использует первую; история сессии содержит обе пары сообщений.")

doc.add_heading("Сценарий 2 Долгосрочная память после нового входа", level=1)
doc.add_paragraph(
    "Клиент явно сообщает имя и цель. Приложение извлекает только эти полезные сведения и записывает их в JSON. "
    "Новый экземпляр приложения имитирует перезапуск сервера."
)
add_dialogue(doc, "Клиент", "Меня зовут Алёна, хочу научиться делать Telegram-ботов")
add_dialogue(doc, "Приложение", "Сохранены поля name и goal")
add_dialogue(doc, "Повторный вход", "Профиль содержит имя «Алёна» и цель «делать Telegram-ботов»")
add_result(doc, "Пройден. Сведения доступны после повторного чтения JSON новым экземпляром приложения.")

doc.add_heading("Сценарий 3 Персонализация двух клиентов", level=1)
doc.add_paragraph(
    "Анна и Борис задают одинаковый вопрос, но получают разные рекомендации. Куратор использует профиль вошедшего пользователя."
)
profiles = doc.add_table(rows=1, cols=3)
profiles.alignment = WD_TABLE_ALIGNMENT.CENTER
profiles.autofit = False
profile_widths = [Inches(1.1), Inches(2.7), Inches(2.7)]
for cell, text, width in zip(profiles.rows[0].cells, ["Поле", "Анна", "Борис"], profile_widths):
    cell.text = text
    cell.width = width
    set_cell_fill(cell, "E8F2EC")
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.paragraphs[0].runs[0].bold = True
for values in [
    ("Цель", "Автоматизировать еженедельные отчёты", "Перейти в Python-разработку"),
    ("Бюджет", "До 12 000 ₽", "До 25 000 ₽"),
    ("Формат", "Самостоятельно", "С наставником"),
    ("Ответ", "Тариф Самостоятельный", "Тариф С наставником"),
]:
    cells = profiles.add_row().cells
    for cell, value, width in zip(cells, values, profile_widths):
        cell.text = value
        cell.width = width
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_result(doc, "Пройден. В ответах упоминается цель только текущего клиента; данные пользователей не смешиваются.")

scenario_4 = doc.add_heading("Сценарий 4 Запрет доступа к чужим данным", level=1)
doc.add_paragraph(
    "Клиент Анна попыталась вызвать две административные операции. Общий слой проверки прав отклонил обе операции до обращения к данным Бориса."
)
access = doc.add_table(rows=1, cols=3)
access.alignment = WD_TABLE_ALIGNMENT.CENTER
access.autofit = False
access_widths = [Inches(3.0), Inches(1.1), Inches(2.4)]
for cell, text, width in zip(access.rows[0].cells, ["Действие", "Ответ", "Состояние данных"], access_widths):
    cell.text = text
    cell.width = width
    set_cell_fill(cell, "185C3C")
    set_cell_border(cell)
    set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
for values in [
    ("Запрос списка всех пользователей", "Отказ", "Список не возвращён"),
    ("Изменение цели пользователя Борис", "Отказ", "Исходная цель сохранена"),
]:
    cells = access.add_row().cells
    for cell, value, width in zip(cells, values, access_widths):
        cell.text = value
        cell.width = width
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
add_result(doc, "Пройден. Клиент не может просматривать или изменять чужие данные и выполнять действия администратора.")

doc.add_heading("Дополнительные проверки", level=1)
items = [
    "JSON не содержит открытых демо-паролей. Сохраняются PBKDF2-хеш и индивидуальная соль.",
    "Поле role нельзя изменить через профиль. Неизвестные поля отклоняются.",
    "JavaScript интерфейса прошёл синтаксическую проверку.",
    "Клиентский и административный экраны проверены в браузере.",
]
for text in items:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Воспроизведение проверки", level=1)
doc.add_paragraph("Из корня проекта выполните команду:")
code = doc.add_paragraph()
code.paragraph_format.left_indent = Inches(0.22)
run = code.add_run("python3 -m unittest discover -s tests -v")
run.font.name = "Aptos Mono"
run.font.size = Pt(10)
run.font.color.rgb = RGBColor(24, 92, 60)
summary = doc.add_paragraph()
summary.add_run("Фактический итог  ").bold = True
summary.add_run("Ran 6 tests — OK")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Домашнее задание Lite  •  Занятие 4")
footer_run.font.size = Pt(8.5)
footer_run.font.color.rgb = RGBColor(105, 115, 108)

doc.save(OUTPUT)
print(OUTPUT)
