from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Результаты тестирования Pro.docx"
DIAGRAM = ROOT / "assets" / "sql-table-relationships.png"
GREEN = RGBColor(24, 92, 60)


def build_relationship_diagram():
    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    canvas = Image.new("RGB", (1800, 1250), "white")
    draw = ImageDraw.Draw(canvas)
    regular = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 29)
    bold = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 31)
    small = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 25)
    relation = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 27)
    boxes = {
        "users": (70, 65, 650, [
            ("id", "TEXT", "PK"), ("username", "TEXT", "UNIQUE"), ("role", "TEXT", ""),
            ("display_name", "TEXT", ""), ("password_salt", "TEXT", ""),
            ("password_hash", "TEXT", ""), ("created_at", "TEXT", ""), ("updated_at", "TEXT", ""),
        ]),
        "memories": (1080, 65, 650, [
            ("id", "INTEGER", "PK"), ("user_id", "TEXT", "FK"), ("key", "TEXT", ""),
            ("value", "TEXT", ""), ("created_at", "TEXT", ""), ("updated_at", "TEXT", ""),
        ]),
        "dialogs": (70, 795, 650, [
            ("id", "TEXT", "PK"), ("user_id", "TEXT", "FK"), ("title", "TEXT", ""),
            ("created_at", "TEXT", ""), ("updated_at", "TEXT", ""),
        ]),
        "messages": (1080, 795, 650, [
            ("id", "INTEGER", "PK"), ("dialog_id", "TEXT", "FK"), ("role", "TEXT", ""),
            ("content", "TEXT", ""), ("created_at", "TEXT", ""),
        ]),
    }
    header_height, row_height = 62, 52
    bounds = {}
    for name, (x, y, width, fields) in boxes.items():
        height = header_height + row_height * len(fields)
        bounds[name] = (x, y, x + width, y + height)
        draw.rounded_rectangle((x, y, x + width, y + height), radius=16, fill="#FFFFFF", outline="#C8D0CA", width=3)
        draw.rounded_rectangle((x, y, x + width, y + header_height), radius=16, fill="#185C3C")
        draw.rectangle((x, y + header_height - 16, x + width, y + header_height), fill="#185C3C")
        draw.text((x + 22, y + 13), name, font=bold, fill="white")
        for index, (field, sql_type, marker) in enumerate(fields):
            top = y + header_height + index * row_height
            if index % 2:
                draw.rectangle((x + 2, top, x + width - 2, top + row_height), fill="#F3F7F4")
            draw.line((x, top, x + width, top), fill="#D9D9D9", width=2)
            draw.text((x + 22, top + 9), field, font=regular, fill="#17211B")
            draw.text((x + 355, top + 10), sql_type, font=small, fill="#667168")
            if marker:
                badge_width = 100 if marker == "UNIQUE" else 56
                draw.rounded_rectangle((x + width - badge_width - 18, top + 9, x + width - 18, top + 42), radius=11, fill="#D9EFE2")
                text_box = draw.textbbox((0, 0), marker, font=small)
                text_width = text_box[2] - text_box[0]
                draw.text((x + width - 18 - badge_width / 2 - text_width / 2, top + 10), marker, font=small, fill="#185C3C")

    def horizontal_arrow(start, end, caption):
        draw.line((*start, *end), fill="#185C3C", width=6)
        ex, ey = end
        sx, sy = start
        triangle = [(ex, ey), (ex - 22, ey - 13), (ex - 22, ey + 13)]
        draw.polygon(triangle, fill="#185C3C")
        draw.text((sx + 18, sy + 10), "1", font=relation, fill="#185C3C")
        draw.text((ex - 42, ey + 10), "N", font=relation, fill="#185C3C")
        box = draw.textbbox((0, 0), caption, font=small)
        tx = (sx + ex - box[2]) / 2
        ty = sy - 43
        draw.rounded_rectangle((tx - 12, ty - 6, tx + box[2] + 12, ty + box[3] + 6), radius=10, fill="white")
        draw.text((tx, ty), caption, font=small, fill="#445048")

    def vertical_arrow(start, end, caption):
        draw.line((*start, *end), fill="#185C3C", width=6)
        ex, ey = end
        sx, sy = start
        draw.polygon([(ex, ey), (ex + 13, ey - 22), (ex - 13, ey - 22)], fill="#185C3C")
        draw.text((sx + 16, sy + 8), "1", font=relation, fill="#185C3C")
        draw.text((ex - 42, ey - 42), "N", font=relation, fill="#185C3C")
        box = draw.textbbox((0, 0), caption, font=small)
        tx = sx - box[2] - 35
        ty = (sy + ey - box[3]) / 2
        draw.rounded_rectangle((tx - 12, ty - 6, tx + box[2] + 12, ty + box[3] + 6), radius=10, fill="white")
        draw.text((tx, ty), caption, font=small, fill="#445048")

    users = bounds["users"]
    memories = bounds["memories"]
    dialogs = bounds["dialogs"]
    messages = bounds["messages"]
    horizontal_arrow((users[2], users[1] + 230), (memories[0], memories[1] + 230), "id → user_id")
    vertical_arrow((users[0] + 325, users[3]), (dialogs[0] + 325, dialogs[1]), "id → user_id")
    horizontal_arrow((dialogs[2], dialogs[1] + 185), (messages[0], messages[1] + 185), "id → dialog_id")
    canvas.save(DIAGRAM, dpi=(220, 220), optimize=True)


def fill(cell, color):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(node)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    group = tc_pr.first_child_found_in("w:tcBorders")
    if group is None:
        group = OxmlElement("w:tcBorders")
        tc_pr.append(group)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), color)
        group.append(node)


def margins(cell, amount=125):
    tc_pr = cell._tc.get_or_add_tcPr()
    group = tc_pr.first_child_found_in("w:tcMar")
    if group is None:
        group = OxmlElement("w:tcMar")
        tc_pr.append(group)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(amount))
        node.set(qn("w:type"), "dxa")
        group.append(node)


def table(doc, headers, widths, rows, dark=True):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    for cell, text, width in zip(result.rows[0].cells, headers, widths):
        cell.text = text
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        fill(cell, "185C3C" if dark else "E8F2EC")
        borders(cell)
        margins(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255) if dark else RGBColor(0, 0, 0)
    for row_index, values in enumerate(rows):
        cells = result.add_row().cells
        for cell, text, width in zip(cells, values, widths):
            cell.text = text
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cell)
            margins(cell)
            if dark and row_index % 2:
                fill(cell, "F3F7F4")
    return result


def dialogue(doc, speaker, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(speaker + "  ")
    run.bold = True
    run.font.color.rgb = GREEN
    p.add_run(text)


def result(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("Результат  ")
    run.bold = True
    run.font.color.rgb = GREEN
    p.add_run(text)


build_relationship_diagram()
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.7)
styles["Normal"].font.color.rgb = RGBColor(23, 33, 27)
styles["Normal"].paragraph_format.space_after = Pt(8)
styles["Normal"].paragraph_format.line_spacing = 1.12
styles["Title"].font.name = "Aptos Display"
styles["Title"].font.size = Pt(28)
styles["Title"].font.bold = True
styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(10)
title_style_ppr = styles["Title"].element.get_or_add_pPr()
style_border = title_style_ppr.find(qn("w:pBdr"))
if style_border is not None:
    title_style_ppr.remove(style_border)
for name, size in (("Heading 1", 18), ("Heading 2", 13)):
    style = styles[name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(13)
    style.paragraph_format.space_after = Pt(7)

title = doc.add_paragraph(style="Title")
title.add_run("Результаты тестирования AI куратора Pro")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
run = subtitle.add_run("SQL память история диалогов и разграничение прав пользователей")
run.font.size = Pt(13)
run.font.color.rgb = RGBColor(86, 98, 90)

meta = table(doc, ["Параметр", "Значение"], [Inches(1.55), Inches(5.05)], [
    ("Приложение", "AI-куратор онлайн-школы Практика"),
    ("Дата проверки", "14 октября 2026 года"),
    ("Хранилище", "SQLite с внешними ключами и индексами"),
], dark=False)
doc.add_paragraph()
intro = doc.add_paragraph()
intro.add_run("Итог проверки. ").bold = True
intro.add_run(
    "Все четыре обязательных сценария Pro выполнены. Общий автоматический прогон Lite и Pro завершился "
    "со статусом OK: 12 тестов из 12. Пользователи, долговременные сведения, диалоги и сообщения сохраняются в SQLite."
)

doc.add_heading("Структура SQL", level=1)
schema = table(doc, ["Таблица", "Назначение", "Связь"], [Inches(1.25), Inches(3.15), Inches(2.2)], [
    ("users", "Учётные записи, роли и хеши паролей", "Родитель памяти и диалогов"),
    ("memories", "Факты и предпочтения пользователя", "user_id ссылается на users"),
    ("dialogs", "Метаданные отдельных диалогов", "user_id ссылается на users"),
    ("messages", "Реплики клиента и куратора", "dialog_id ссылается на dialogs"),
])

diagram_break = doc.add_paragraph()
diagram_break.add_run().add_break(WD_BREAK.PAGE)
diagram_spacer = doc.add_paragraph()
diagram_spacer.paragraph_format.space_after = Pt(10)
doc.add_heading("Диаграмма связей таблиц", level=1)
diagram_intro = doc.add_paragraph(
    "На диаграмме показаны поля каждой таблицы и связи внешних ключей. Обозначения 1 и N соответствуют связи один ко многим."
)
diagram_intro.paragraph_format.space_after = Pt(10)
diagram_paragraph = doc.add_paragraph()
diagram_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
shape = diagram_paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(6.65))
doc_pr = shape._inline.docPr
doc_pr.set("descr", "Диаграмма таблиц users memories dialogs messages с полями и связями один ко многим")
caption = doc.add_paragraph("Связи SQL таблиц приложения Pro")
caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
caption.paragraph_format.space_before = Pt(5)
caption.paragraph_format.space_after = Pt(0)
for caption_run in caption.runs:
    caption_run.italic = True
    caption_run.font.size = Pt(9.5)
    caption_run.font.color.rgb = RGBColor(86, 98, 90)

scenario_1 = doc.add_heading("Сценарий 1 Память текущего диалога", level=1)
scenario_1.paragraph_format.page_break_before = True
doc.add_paragraph("Куратор получает последние сообщения только выбранного диалога. Сообщения сохраняются в SQL в порядке добавления.")
dialogue(doc, "Клиент", "Хочу разобраться с функциями")
dialogue(doc, "Клиент", "Продолжи рекомендацию")
dialogue(doc, "Куратор", "Продолжаю текущий диалог и помню, что ранее вы написали: «Хочу разобраться с функциями»")
result(doc, "Пройден. В таблице messages сохранены две реплики клиента и два ответа куратора.")

doc.add_heading("Сценарий 2 Долгосрочная память между диалогами", level=1)
doc.add_paragraph(
    "В первом диалоге клиент сообщил цель. Затем созданы второй диалог и новый экземпляр приложения, имитирующий перезапуск сервера."
)
dialogue(doc, "Клиент", "Моя цель — научиться анализировать продажи")
dialogue(doc, "Второй диалог", "На вопрос о памяти куратор назвал сохранённую цель")
dialogue(doc, "После перезапуска", "Третий диалог также использовал цель из таблицы memories")
result(doc, "Пройден. Долгосрочное сведение не зависит от текущего диалога или процесса приложения.")

doc.add_heading("Сценарий 3 Изменение базы влияет на ответ", level=1)
changes = table(doc, ["Состояние", "Формат в SQL", "Ответ куратора"], [Inches(1.35), Inches(2.45), Inches(2.8)], [
    ("До изменения", "самостоятельно", "Тариф Самостоятельный"),
    ("После изменения", "с наставником", "Тариф С наставником"),
], dark=False)
result(doc, "Пройден. Новый диалог использует актуальное значение из SQL, а не старую копию профиля.")

page_break = doc.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)
scenario_4 = doc.add_heading("Сценарий 4 Запрет доступа к чужим данным", level=1)
doc.add_paragraph("Клиент Анна выполнила четыре действия, направленные на данные Бориса. Проверка владельца отклонила каждое действие до чтения или изменения данных.")
access = table(doc, ["Действие клиента", "Ответ", "Состояние данных"], [Inches(3.0), Inches(1.0), Inches(2.6)], [
    ("Получить список всех пользователей", "Отказ", "Список не возвращён"),
    ("Изменить память Бориса", "Отказ", "Память не изменена"),
    ("Удалить память Бориса", "Отказ", "Память сохранена"),
    ("Открыть диалог Бориса", "Отказ", "Сообщения не возвращены"),
])
result(doc, "Пройден. Клиент не может читать или изменять чужую память и историю диалогов.")

doc.add_heading("Дополнительные проверки", level=1)
for item in [
    "Клиент может просмотреть, изменить и удалить собственное сохранённое сведение.",
    "Внешний ключ messages.dialog_id ссылается на dialogs.id.",
    "Каждое соединение завершается commit или rollback и гарантированно закрывается.",
    "Клиентский и административный экраны Pro проверены в браузере.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Воспроизведение проверки", level=1)
doc.add_paragraph("Из корня проекта выполните команду:")
code = doc.add_paragraph()
code.paragraph_format.left_indent = Inches(0.2)
run = code.add_run("PYTHONWARNINGS=error::ResourceWarning python3 -m unittest discover -s tests -v")
run.font.name = "Aptos Mono"
run.font.size = Pt(9.5)
run.font.color.rgb = GREEN
summary = doc.add_paragraph()
summary.add_run("Фактический итог  ").bold = True
summary.add_run("Ran 12 tests — OK")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Домашнее задание Pro  •  Занятие 4")
footer_run.font.size = Pt(8.5)
footer_run.font.color.rgb = RGBColor(105, 115, 108)

doc.save(OUTPUT)
print(OUTPUT)
